from io import BytesIO
from typing import Dict, Any
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.drawing.image import Image as OpenPyXLImage

from core.chart_exporter import (
    generate_price_technical_chart_img,
    generate_radar_chart_img,
    generate_correlation_heatmap_img,
    generate_portfolio_growth_chart_img,
    generate_sector_bar_chart_img
)


def generate_excel_bytes(session_data: Dict[str, Any]) -> bytes:
    wb = openpyxl.Workbook()

    font_family = "Segoe UI"
    font_title = Font(name=font_family, size=14, bold=True, color="0F172A")
    font_section = Font(name=font_family, size=11, bold=True, color="0284C7")
    font_header = Font(name=font_family, size=10, bold=True, color="0F172A")
    font_data = Font(name=font_family, size=10, color="0F172A")

    fill_header = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")

    border_thin = Border(
        left=Side(style="thin", color="E2E8F0"),
        right=Side(style="thin", color="E2E8F0"),
        top=Side(style="thin", color="E2E8F0"),
        bottom=Side(style="thin", color="E2E8F0")
    )

    ticker = session_data.get("ticker", "ASSET")
    company = session_data.get("company_name", ticker)
    curr = session_data.get("currency_symbol", "₹")

    # ==========================================
    # TAB 1: Executive Summary
    # ==========================================
    ws_summary = wb.active
    ws_summary.title = "Executive Summary"
    ws_summary.views.sheetView[0].showGridLines = True

    ws_summary["A1"] = f"THE EXECUTIVE TERMINAL — {ticker} ({company})"
    ws_summary["A1"].font = font_title

    ws_summary["A3"] = "KEY PERFORMANCE INDICATORS"
    ws_summary["A3"].font = font_section

    dhs = session_data.get("dhs", {})
    risk = session_data.get("risk_metrics", {})

    kpi_headers = ["Metric", "Value"]
    kpi_rows = [
        ["Data Health Score (DHS)", f"{dhs.get('dhs_score')}% ({dhs.get('status')})"],
        ["Benchmark Index Used", session_data.get("benchmark_symbol", "^NSEI")],
        ["Total Capital Invested", f"{curr}{risk.get('total_invested', 0.0):,}"],
        ["Current Market Value", f"{curr}{risk.get('current_value', 0.0):,}"],
        ["Total Unrealized P&L", f"{curr}{risk.get('total_pnl', 0.0):,} ({risk.get('pnl_pct', 0.0)}%)"],
        ["Portfolio Beta (Systematic Risk)", risk.get("portfolio_beta", 1.0)],
        ["Annualized Sharpe Ratio", risk.get("sharpe_ratio", 0.0)],
    ]

    for col_idx, text in enumerate(kpi_headers, start=1):
        cell = ws_summary.cell(row=4, column=col_idx, value=text)
        cell.font = font_header
        cell.fill = fill_header
        cell.border = border_thin

    for row_idx, r_data in enumerate(kpi_rows, start=5):
        for col_idx, val in enumerate(r_data, start=1):
            cell = ws_summary.cell(row=row_idx, column=col_idx, value=val)
            cell.font = font_data
            cell.border = border_thin

    # Decision Support Bullets
    start_r = 14
    ws_summary.cell(row=start_r, column=1, value="AUTOMATED DECISION SUPPORT TAKEAWAYS").font = font_section

    insights = session_data.get("insights", [])
    for idx, bullet in enumerate(insights, start=start_r + 1):
        clean_b = bullet.replace("**", "").replace("**", "")
        cell = ws_summary.cell(row=idx, column=1, value=f"• {clean_b}")
        cell.font = font_data

    # Embed Radar & Growth Charts in Tab 1
    radar_img_buf = generate_radar_chart_img(session_data.get("radar_scores", {}), ticker=ticker)
    if radar_img_buf:
        img_radar = OpenPyXLImage(radar_img_buf)
        img_radar.width = 380
        img_radar.height = 320
        ws_summary.add_image(img_radar, "D3")

    bench_symbol = session_data.get("benchmark_symbol", "^NSEI")
    growth_img_buf = generate_portfolio_growth_chart_img(bench_name=bench_symbol)
    if growth_img_buf:
        img_growth = OpenPyXLImage(growth_img_buf)
        img_growth.width = 440
        img_growth.height = 280
        ws_summary.add_image(img_growth, "D20")

    # ==========================================
    # TAB 2: Stock Analytics
    # ==========================================
    ws_stock = wb.create_sheet(title="Stock Analytics")
    ws_stock.views.sheetView[0].showGridLines = True

    ws_stock["A1"] = f"TECHNICAL & FUNDAMENTAL ANALYSIS — {ticker}"
    ws_stock["A1"].font = font_title

    ohlcv_data = session_data.get("ohlcv", [])
    if ohlcv_data:
        df_ohlcv = pd.DataFrame(ohlcv_data)
        cols = [c for c in ["Date", "Open", "High", "Low", "Close", "Volume", "EMA_20", "EMA_50", "EMA_200", "RSI_14", "MACD"] if c in df_ohlcv.columns]

        for col_idx, col_name in enumerate(cols, start=1):
            cell = ws_stock.cell(row=3, column=col_idx, value=col_name)
            cell.font = font_header
            cell.fill = fill_header
            cell.border = border_thin

        for row_idx, row_dict in enumerate(df_ohlcv[cols].to_dict("records"), start=4):
            for col_idx, col_name in enumerate(cols, start=1):
                val = row_dict.get(col_name)
                cell = ws_stock.cell(row=row_idx, column=col_idx, value=str(val))
                cell.font = font_data
                cell.border = border_thin

    # Embed Price Technical Chart in Tab 2
    price_img_buf = generate_price_technical_chart_img(ohlcv_data, ticker=ticker)
    if price_img_buf:
        img_price = OpenPyXLImage(price_img_buf)
        img_price.width = 600
        img_price.height = 340
        ws_stock.add_image(img_price, "M3")

    # ==========================================
    # TAB 3: Portfolio Terminal
    # ==========================================
    ws_port = wb.create_sheet(title="Portfolio Holdings")
    ws_port.views.sheetView[0].showGridLines = True

    ws_port["A1"] = "PORTFOLIO HOLDINGS & CORRELATION"
    ws_port["A1"].font = font_title

    holdings = risk.get("summary_table", [])
    if holdings:
        df_hold = pd.DataFrame(holdings)
        h_cols = list(df_hold.columns)

        for col_idx, col_name in enumerate(h_cols, start=1):
            cell = ws_port.cell(row=3, column=col_idx, value=col_name)
            cell.font = font_header
            cell.fill = fill_header
            cell.border = border_thin

        for row_idx, row_dict in enumerate(df_hold.to_dict("records"), start=4):
            for col_idx, col_name in enumerate(h_cols, start=1):
                val = row_dict.get(col_name)
                cell = ws_port.cell(row=row_idx, column=col_idx, value=val)
                cell.font = font_data
                cell.border = border_thin

    # Embed Correlation Heatmap in Tab 3
    corr_img_buf = generate_correlation_heatmap_img(risk.get("correlation_matrix"))
    if corr_img_buf:
        img_corr = OpenPyXLImage(corr_img_buf)
        img_corr.width = 420
        img_corr.height = 330
        ws_port.add_image(img_corr, "J3")

    # ==========================================
    # TAB 4: Pipeline Telemetry & Sector Allocation
    # ==========================================
    ws_telem = wb.create_sheet(title="Pipeline Telemetry")
    ws_telem.views.sheetView[0].showGridLines = True

    ws_telem["A1"] = "DATA HEALTH INDEX & SANITIZATION AUDIT LOG"
    ws_telem["A1"].font = font_title

    ws_telem["A3"] = "Imputation & Sanitization Events"
    ws_telem["A3"].font = font_section

    imputations = session_data.get("imputed_log", [])
    ws_telem.cell(row=4, column=1, value="Sanitization Event Description").font = font_header
    ws_telem.cell(row=4, column=1).fill = fill_header
    ws_telem.cell(row=4, column=1).border = border_thin

    if imputations:
        for r_i, imp_text in enumerate(imputations, start=5):
            c = ws_telem.cell(row=r_i, column=1, value=imp_text)
            c.font = font_data
            c.border = border_thin
    else:
        c = ws_telem.cell(row=5, column=1, value="No imputation required. Payload 100% complete.")
        c.font = font_data
        c.border = border_thin

    # Embed Sector Bar Chart in Tab 4
    treemap_payload = session_data.get("market_radar", {}).get("treemap", [])
    sector_img_buf = generate_sector_bar_chart_img(treemap_payload)
    if sector_img_buf:
        img_sector = OpenPyXLImage(sector_img_buf)
        img_sector.width = 460
        img_sector.height = 280
        ws_telem.add_image(img_sector, "C3")

    # Auto-adjust column widths across all worksheets
    for sheet in wb.worksheets:
        for col in sheet.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = get_column_letter(col[0].column)
            sheet.column_dimensions[col_letter].width = max(max_len + 3, 12)

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()